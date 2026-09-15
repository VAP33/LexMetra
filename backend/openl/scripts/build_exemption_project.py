"""Generate the OpenL rule project for LMPC Rule 3 scope/exemption + Rule 26(a).

This re-expresses the *legal decision* currently implemented in
`backend/exemption.py` (`classify_exemption`) as an OpenL **decision table**.

Design (see docs/cdd/agents/RULE-01/DECISIONS.md):
- The OpenL table owns the LEGAL content: the branch precedence, the Rule 3
  quantity thresholds, the cement/fertilizer exception, and the Rule 26(a)
  small-pack thresholds + excluded categories.
- Thresholds are READ FROM rules/rules.json here and baked into the table, so
  rules.json stays the single source of truth (no hand-retyped constants).
- Unit normalization (g/kg/ml/l, weight!=volume) and "quantity established"
  validation stay in Python (they are data-cleaning, not legal policy). Python
  passes OpenL canonical inputs: massG / volumeMl as doubles (-1.0 sentinel =
  "not applicable"), plus quantityEstablished + the tri-state booleans.
- The table returns the `exemption_type` string (the legal classification). The
  Python resolver maps that string to the full ExemptionResult/RuleFinding
  fields (is_exempt, rule_id, review_required, reason), which are deterministic
  consequences of the type, not independent legal thresholds.

Output: backend/openl/dist/lmpc-exemption.zip
  - rules-deploy.xml        service name + RESTFUL publisher
  - LmpcExemption.xlsx      one decision table: String classifyExemption(...)
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
RULES_PATH = REPO / "rules" / "rules.json"
DIST = HERE.parent / "dist"
DIST.mkdir(parents=True, exist_ok=True)

SERVICE_NAME = "lmpc_exemption"
DEPLOYMENT_NAME = "lexmetra"

RULES_DEPLOY_XML = f"""<?xml version="1.0" encoding="UTF-8"?>
<rules-deploy>
    <serviceName>{SERVICE_NAME}</serviceName>
    <provideRuntimeContext>false</provideRuntimeContext>
    <publishers>
        <publisher>RESTFUL</publisher>
    </publishers>
</rules-deploy>
"""

# Excluded-category sets, transcribed to match backend/exemption.py verbatim.
CEMENT_SET = "cement,fertilizer,agricultural farm produce"
TOBACCO_SET = "tobacco,tobacco product,tobacco products,pan masala"
KNOWN_RETAIL_TYPES = "retail,ecommerce,export"


def load_thresholds() -> dict:
    data = json.loads(RULES_PATH.read_text(encoding="utf-8"))
    rules = {r["rule_id"]: r for r in data["rules"]}
    r3 = rules["LMPC-2011-R3-SCOPE"]["threshold"]
    r26 = rules["LMPC-2011-R26-SMALL-PACKS"]["threshold"]
    general = r3["general_package_quantity_exemption"]
    cement = r3["cement_fertilizer_exception"]
    upper = r26["small_pack_upper_bound"]
    return {
        # Rule 3 general weight exclusion (kg -> g)
        "mass_gt_general_g": float(general["weight_kg_gt"]) * 1000.0,
        # Rule 3 cement/fertilizer weight exclusion (kg -> g)
        "mass_gt_cement_g": float(cement["weight_kg_gt"]) * 1000.0,
        # Rule 3 general volume exclusion (l -> ml)
        "vol_gt_general_ml": float(general["volume_l_gt"]) * 1000.0,
        # Rule 26(a) small-pack upper bounds
        "mass_le_small_g": float(upper["weight_g_lte"]),
        "vol_le_small_ml": float(upper["volume_ml_lte"]),
    }


# Column plan (Excel A..). Each condition column carries: id, expr, decl, title.
# An empty data cell means "condition not applied for this rule".
COLUMNS = [
    ("C1", "contains(saleTypes, saleType)", "String[] saleTypes", "Sale type in"),
    ("C2", "contains(categoriesIn, productCategory)", "String[] categoriesIn", "Category in"),
    ("C3", "!contains(categoriesNotIn, productCategory)", "String[] categoriesNotIn", "Category not in"),
    ("C4", "isPrepackagedFalse == prepackagedFalseFlag", "Boolean prepackagedFalseFlag", "Prepackaged=false"),
    ("C5", "directIndInst == directFlag", "Boolean directFlag", "Direct ind/inst"),
    ("C6", "isExportOnly == exportFlag", "Boolean exportFlag", "Export only"),
    ("C7", "quantityEstablished == qtyEstFlag", "Boolean qtyEstFlag", "Qty established"),
    ("C8", "massG > massGt", "Double massGt", "Mass g >"),
    ("C9", "massG >= 0 && massG <= massLe", "Double massLe", "Mass g <="),
    ("C10", "volumeMl > volGt", "Double volGt", "Vol ml >"),
    ("C11", "volumeMl >= 0 && volumeMl <= volLe", "Double volLe", "Vol ml <="),
    ("C12", "!contains(saleTypesNotIn, saleType)", "String[] saleTypesNotIn", "Sale type not in"),
    ("RET1", "ret", "String ret", "exemption_type"),
]

METHOD_SIG = (
    "Rules String classifyExemption("
    "String saleType, String productCategory, double massG, double volumeMl, "
    "boolean quantityEstablished, boolean isPrepackagedFalse, "
    "boolean directIndInst, boolean isExportOnly)"
)


def build_rows(t: dict) -> list[dict]:
    """Ordered decision rules. Keys are column indices 0..11 (conditions) + 'ret'.

    Order mirrors backend/exemption.py classify_exemption() precedence exactly.
    """
    T, F = True, False
    return [
        # 1. explicitly not pre-packaged
        {3: T, "ret": "not_prepackaged"},
        # 2. explicit direct industrial/institutional
        {4: T, "ret": "industrial_or_institutional_direct_sale"},
        # 3. industrial/institutional sale type
        {0: "industrial,institutional", "ret": "industrial_or_institutional_direct_sale"},
        # 4. export-only AND export sale
        {0: "export", 5: T, "ret": "export_only_transaction"},
        # 5. net quantity not established
        {6: F, "ret": "quantity_not_established"},
        # 6. Rule 3 weight exclusion, cement/fertilizer/agri (> 50 kg)
        {1: CEMENT_SET, 7: t["mass_gt_cement_g"], "ret": "rule_3_quantity_exclusion"},
        # 7. Rule 3 weight exclusion, all other categories (> 25 kg)
        {2: CEMENT_SET, 7: t["mass_gt_general_g"], "ret": "rule_3_quantity_exclusion"},
        # 8. Rule 3 volume exclusion (> 25 l)
        {9: t["vol_gt_general_ml"], "ret": "rule_3_quantity_exclusion"},
        # 9a. Rule 26(a) small pack by mass (<= 10 g), not an excluded category
        {2: TOBACCO_SET, 8: t["mass_le_small_g"], "ret": "rule_26_small_pack"},
        # 9b. Rule 26(a) small pack by volume (<= 10 ml), not an excluded category
        {2: TOBACCO_SET, 10: t["vol_le_small_ml"], "ret": "rule_26_small_pack"},
        # 10. wholesale = different declaration regime (not an exemption)
        {0: "wholesale", "ret": "different_declaration_regime"},
        # 11. unrecognized sale type
        {11: KNOWN_RETAIL_TYPES, "ret": "unknown_sale_type"},
        # 12. default: not exempt, no specific type
        {"ret": "none"},
    ]


def _cell(v) -> object:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return v


def build_xlsx(t: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"

    ncols = len(COLUMNS)
    # Row 1: header signature, merged across all columns.
    ws.cell(row=1, column=1, value=METHOD_SIG)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    # Rows 2-5: id / expr / decl / title
    for ci, (cid, expr, decl, title) in enumerate(COLUMNS, start=1):
        ws.cell(row=2, column=ci, value=cid)
        ws.cell(row=3, column=ci, value=expr)
        ws.cell(row=4, column=ci, value=decl)
        ws.cell(row=5, column=ci, value=title)
    # Rows 6+: data
    rows = build_rows(t)
    for ri, row in enumerate(rows, start=6):
        for ci in range(len(COLUMNS) - 1):  # condition columns
            if ci in row:
                ws.cell(row=ri, column=ci + 1, value=_cell(row[ci]))
        ws.cell(row=ri, column=ncols, value=row["ret"])  # RET column (last)

    # Widen a bit for human inspection.
    for ci in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    t = load_thresholds()
    print("Thresholds from rules.json:", json.dumps(t))
    xlsx = build_xlsx(t)
    out = DIST / "lmpc-exemption.zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("rules-deploy.xml", RULES_DEPLOY_XML)
        z.writestr("LmpcExemption.xlsx", xlsx)
    # Also drop the raw xlsx next to the zip for inspection / version control.
    (DIST / "LmpcExemption.xlsx").write_bytes(xlsx)
    print(f"wrote {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
