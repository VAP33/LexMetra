"""Build a trivial OpenL project zip to de-risk the deploy->execute pipeline.

Produces backend/openl/dist/smoke.zip containing:
  - rules-deploy.xml  (service name + RESTFUL publisher, no runtime context)
  - Rules.xlsx        (one SimpleRules table: String hello(String name))
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
DIST = HERE.parent / "dist"
DIST.mkdir(parents=True, exist_ok=True)

RULES_DEPLOY_XML = """<?xml version="1.0" encoding="UTF-8"?>
<rules-deploy>
    <serviceName>smoke</serviceName>
    <provideRuntimeContext>false</provideRuntimeContext>
    <publishers>
        <publisher>RESTFUL</publisher>
    </publishers>
</rules-deploy>
"""


def build_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    # SimpleRules: first column = condition on the single param, last = return.
    ws["A1"] = "SimpleRules String hello(String name)"
    ws["A2"] = "Name"
    ws["B2"] = "Greeting"
    ws["A3"] = "World"
    ws["B3"] = "Hi, World!"
    ws["A4"] = None  # empty condition => matches anything (default row)
    ws["B4"] = "Hi, stranger!"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    xlsx = build_xlsx()
    out = DIST / "smoke.zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("rules-deploy.xml", RULES_DEPLOY_XML)
        z.writestr("Rules.xlsx", xlsx)
    print(f"wrote {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
